"""
Server-side Python statement generation and financial transaction export service for BankVCS 2.0.
Supports E-Passbook, Mini Statements, PDF/CSV/XLSX exports, and Transaction Receipts.
"""

from datetime import datetime, date, timedelta
from decimal import Decimal
from typing import List, Dict, Any, Optional
import io
import csv

from app import db
from app.models.transaction import Transaction, TransactionType, TransactionStatus, TransactionMode
from app.models.account import Account


def resolve_date_preset(preset: Optional[str], start_date: Optional[date] = None, end_date: Optional[date] = None):
    """Resolves date filter presets into start_date and end_date."""
    today = date.today()
    if preset == "7_DAYS":
        return today - timedelta(days=7), today
    elif preset == "30_DAYS":
        return today - timedelta(days=30), today
    elif preset == "3_MONTHS":
        return today - timedelta(days=90), today
    elif preset == "CURRENT_FY":
        if today.month >= 4:
            fy_start = date(today.year, 4, 1)
        else:
            fy_start = date(today.year - 1, 4, 1)
        return fy_start, today
    elif preset == "PREVIOUS_FY":
        if today.month >= 4:
            fy_start = date(today.year - 1, 4, 1)
            fy_end = date(today.year, 3, 31)
        else:
            fy_start = date(today.year - 2, 4, 1)
            fy_end = date(today.year - 1, 3, 31)
        return fy_start, fy_end
    return start_date, end_date


def mask_account_number(acc_num: str) -> str:
    """Masks account number keeping only last 4 digits visible."""
    if not acc_num or len(acc_num) < 4:
        return acc_num or "N/A"
    return f"XXXX XXXX {acc_num[-4:]}"


def generate_account_statement(
    account_id: int,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    transaction_type: Optional[str] = None,
    status: Optional[str] = None,
    mode: Optional[str] = None,
    search: Optional[str] = None,
    preset: Optional[str] = None,
    sort_order: str = "desc"
) -> Dict[str, Any]:
    """
    Generates server-side financial statement summary and transaction list for an account.
    """
    account = db.session.get(Account, account_id)
    if not account:
        return {"success": False, "error": "Account not found"}

    start_date, end_date = resolve_date_preset(preset, start_date, end_date)

    query = Transaction.query.filter_by(account_id=account_id)

    if start_date:
        query = query.filter(Transaction.created_at >= datetime.combine(start_date, datetime.min.time()))
    if end_date:
        query = query.filter(Transaction.created_at <= datetime.combine(end_date, datetime.max.time()))

    if transaction_type and transaction_type != "ALL":
        query = query.filter_by(transaction_type=transaction_type)

    if status and status != "ALL":
        query = query.filter_by(status=status)

    if mode and mode != "ALL":
        query = query.filter_by(transaction_mode=mode)

    if search:
        search_term = f"%{search.strip()}%"
        query = query.filter(
            (Transaction.description.ilike(search_term)) |
            (Transaction.reference_number.ilike(search_term))
        )

    # Fetch chronologically for ledger reconciliation
    all_filtered_txs = query.order_by(Transaction.created_at.asc(), Transaction.id.asc()).all()

    # Calculate Opening Balance (if start_date is set)
    opening_balance = Decimal("0.00")
    if start_date:
        prior_txs = Transaction.query.filter(
            Transaction.account_id == account_id,
            Transaction.status == TransactionStatus.COMPLETED,
            Transaction.created_at < datetime.combine(start_date, datetime.min.time())
        ).order_by(Transaction.created_at.asc(), Transaction.id.asc()).all()

        running_prev = Decimal("0.00")
        for ptx in prior_txs:
            amt = Decimal(ptx.amount)
            if ptx.balance_after is not None:
                if Decimal(ptx.balance_after) >= running_prev:
                    running_prev = Decimal(ptx.balance_after)
                else:
                    running_prev = Decimal(ptx.balance_after)
            else:
                if ptx.transaction_type in [TransactionType.DEPOSIT, TransactionType.REVERSAL]:
                    running_prev += amt
                else:
                    running_prev -= amt
        opening_balance = running_prev

    total_credits = Decimal("0.00")
    total_debits = Decimal("0.00")
    running = opening_balance

    for tx in all_filtered_txs:
        if tx.status != TransactionStatus.COMPLETED:
            continue
        amt = Decimal(tx.amount)
        if tx.balance_after is not None:
            bal = Decimal(tx.balance_after)
            if bal >= running:
                total_credits += (bal - running)
            else:
                total_debits += (running - bal)
            running = bal
        else:
            if tx.transaction_type in [TransactionType.DEPOSIT, TransactionType.REVERSAL]:
                total_credits += amt
                running += amt
            else:
                total_debits += amt
                running -= amt

    if start_date:
        closing_balance = running
    else:
        closing_balance = Decimal(account.balance)

    # Sort final return list according to sort_order parameter
    if sort_order.lower() == "desc":
        return_txs = sorted(all_filtered_txs, key=lambda t: (t.created_at, t.id), reverse=True)
    else:
        return_txs = all_filtered_txs

    owner_name = account.owner.full_name or account.owner.username if account.owner else "Valued Customer"

    return {
        "success": True,
        "account_id": account.id,
        "account_number": account.account_number,
        "masked_account_number": mask_account_number(account.account_number),
        "account_type": account.account_type,
        "customer_id": account.user_id,
        "customer_name": owner_name,
        "current_balance": float(account.balance),
        "available_balance": float(account.available_balance),
        "opening_balance": float(opening_balance),
        "closing_balance": float(closing_balance),
        "total_credits": float(total_credits),
        "total_debits": float(total_debits),
        "transaction_count": len(return_txs),
        "start_date": start_date.strftime("%Y-%m-%d") if start_date else "Account Inception",
        "end_date": end_date.strftime("%Y-%m-%d") if end_date else date.today().strftime("%Y-%m-%d"),
        "generated_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"),
        "transactions": return_txs,
    }


def export_statement_csv(
    account_id: int,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    transaction_type: Optional[str] = None,
    status: Optional[str] = None,
    mode: Optional[str] = None,
    search: Optional[str] = None,
    preset: Optional[str] = None
) -> str:
    """Generates CSV format statement string for download."""
    data = generate_account_statement(
        account_id, start_date, end_date, transaction_type, status, mode, search, preset
    )
    if not data.get("success"):
        return ""

    output = io.StringIO()
    writer = csv.writer(output)

    writer.writerow(["BANKVCS 2.0 OFFICIAL E-PASSBOOK / STATEMENT"])
    writer.writerow(["Account Holder", data["customer_name"]])
    writer.writerow(["Account Number", data["masked_account_number"]])
    writer.writerow(["Account Type", data["account_type"]])
    writer.writerow(["Statement Period", f"{data['start_date']} to {data['end_date']}"])
    writer.writerow(["Generated Date", data["generated_at"]])
    writer.writerow([])
    writer.writerow(["FINANCIAL SUMMARY"])
    writer.writerow(["Opening Balance", f"INR {data['opening_balance']:.2f}"])
    writer.writerow(["Total Credits", f"INR {data['total_credits']:.2f}"])
    writer.writerow(["Total Debits", f"INR {data['total_debits']:.2f}"])
    writer.writerow(["Closing Balance", f"INR {data['closing_balance']:.2f}"])
    writer.writerow([])
    writer.writerow(["Date", "Reference Number", "Narration", "Mode", "Type", "Amount", "Credit (INR)", "Debit (INR)", "Balance After (INR)", "Status", "Description"])

    running = Decimal(data["opening_balance"])
    for tx in data["transactions"]:
        amt = Decimal(tx.amount)
        credit = "0.00"
        debit = "0.00"
        if tx.balance_after is not None:
            bal = Decimal(tx.balance_after)
            if bal >= running:
                credit = f"{float(bal - running):.2f}"
            else:
                debit = f"{float(running - bal):.2f}"
            running = bal
        else:
            if tx.transaction_type in ["DEPOSIT", "REVERSAL"]:
                credit = f"{float(amt):.2f}"
            else:
                debit = f"{float(amt):.2f}"

        writer.writerow([
            tx.created_at.strftime("%Y-%m-%d %H:%M:%S") if tx.created_at else "",
            tx.reference_number,
            tx.narration,
            tx.transaction_mode or "TRANSFER",
            tx.transaction_type,
            f"{float(amt):.2f}",
            credit,
            debit,
            tx.balance_after_transaction,
            tx.status,
            tx.description or ""
        ])

    return output.getvalue()


def export_statement_xlsx(
    account_id: int,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    transaction_type: Optional[str] = None,
    status: Optional[str] = None,
    mode: Optional[str] = None,
    search: Optional[str] = None,
    preset: Optional[str] = None
) -> bytes:
    """Generates Excel (.xlsx) statement byte buffer."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    data = generate_account_statement(
        account_id, start_date, end_date, transaction_type, status, mode, search, preset
    )
    if not data.get("success"):
        return b""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "E-Passbook Statement"

    # Styling
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Arial", size=14, bold=True, color="0F172A")
    sub_font = Font(name="Arial", size=10, italic=True, color="64748B")
    bold_font = Font(name="Arial", size=10, bold=True)

    ws.append(["BankVCS 2.0 - Digital Banking E-Passbook Statement"])
    ws.cell(row=1, column=1).font = title_font
    ws.append([f"Generated on {data['generated_at']}"])
    ws.cell(row=2, column=1).font = sub_font
    ws.append([])

    # Account Metadata
    meta_rows = [
        ("Account Holder", data["customer_name"]),
        ("Account Number", data["masked_account_number"]),
        ("Account Type", data["account_type"]),
        ("Statement Period", f"{data['start_date']} to {data['end_date']}"),
        ("Opening Balance", f"INR {data['opening_balance']:,.2f}"),
        ("Total Credits (+)", f"INR {data['total_credits']:,.2f}"),
        ("Total Debits (-)", f"INR {data['total_debits']:,.2f}"),
        ("Closing Balance", f"INR {data['closing_balance']:,.2f}"),
    ]
    for label, val in meta_rows:
        ws.append([label, val])
        r = ws.max_row
        ws.cell(row=r, column=1).font = bold_font

    ws.append([])

    # Table Header
    headers = ["Date & Time", "Reference No.", "Description / Narration", "Mode", "Type", "Credit (INR)", "Debit (INR)", "Balance After (INR)", "Status"]
    ws.append(headers)
    header_row_idx = ws.max_row
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    running = Decimal(data["opening_balance"])
    for tx in data["transactions"]:
        amt = Decimal(tx.amount)
        credit = 0.0
        debit = 0.0
        if tx.balance_after is not None:
            bal = Decimal(tx.balance_after)
            if bal >= running:
                credit = float(bal - running)
            else:
                debit = float(running - bal)
            running = bal
        else:
            if tx.transaction_type in ["DEPOSIT", "REVERSAL"]:
                credit = float(amt)
            else:
                debit = float(amt)

        ws.append([
            tx.created_at.strftime("%Y-%m-%d %H:%M:%S") if tx.created_at else "",
            tx.reference_number,
            tx.narration,
            tx.transaction_mode or "TRANSFER",
            tx.transaction_type,
            credit if credit > 0 else "-",
            debit if debit > 0 else "-",
            float(tx.balance_after or 0.0),
            tx.status
        ])

    # Auto-adjust column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_statement_pdf(
    account_id: int,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    transaction_type: Optional[str] = None,
    status: Optional[str] = None,
    mode: Optional[str] = None,
    search: Optional[str] = None,
    preset: Optional[str] = None
) -> bytes:
    """Generates professional PDF E-Passbook Bank Statement byte buffer."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    data = generate_account_statement(
        account_id, start_date, end_date, transaction_type, status, mode, search, preset
    )
    if not data.get("success"):
        return b""

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0F172A")
    )
    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#64748B")
    )
    header_style = ParagraphStyle(
        'HeaderStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1
    )
    cell_style = ParagraphStyle(
        'CellStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1E293B")
    )

    elements = []

    # Title Banner
    elements.append(Paragraph("<b>BankVCS 2.0</b> &bull; Digital Banking E-Passbook", title_style))
    elements.append(Paragraph(f"Generated on {data['generated_at']} | Official Account Statement", sub_style))
    elements.append(Spacer(1, 10))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#2563EB"), spaceAfter=12))

    # Account Metadata & Summary Grid
    meta_data = [
        [
            Paragraph(f"<b>Account Holder:</b> {data['customer_name']}", cell_style),
            Paragraph(f"<b>Opening Balance:</b> INR {data['opening_balance']:,.2f}", cell_style)
        ],
        [
            Paragraph(f"<b>Account Number:</b> {data['masked_account_number']}", cell_style),
            Paragraph(f"<b>Total Credits (+):</b> INR {data['total_credits']:,.2f}", cell_style)
        ],
        [
            Paragraph(f"<b>Account Type:</b> {data['account_type']}", cell_style),
            Paragraph(f"<b>Total Debits (-):</b> INR {data['total_debits']:,.2f}", cell_style)
        ],
        [
            Paragraph(f"<b>Statement Period:</b> {data['start_date']} to {data['end_date']}", cell_style),
            Paragraph(f"<b>Closing Balance:</b> INR {data['closing_balance']:,.2f}", cell_style)
        ]
    ]

    meta_table = Table(meta_data, colWidths=[270, 270])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#E2E8F0")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    elements.append(meta_table)
    elements.append(Spacer(1, 15))

    # Transaction Table Header
    tx_table_data = [
        [
            Paragraph("<b>Date &amp; Time</b>", header_style),
            Paragraph("<b>Reference No.</b>", header_style),
            Paragraph("<b>Narration</b>", header_style),
            Paragraph("<b>Mode</b>", header_style),
            Paragraph("<b>Credit</b>", header_style),
            Paragraph("<b>Debit</b>", header_style),
            Paragraph("<b>Balance</b>", header_style),
            Paragraph("<b>Status</b>", header_style)
        ]
    ]

    running = Decimal(data["opening_balance"])
    for tx in data["transactions"]:
        amt = Decimal(tx.amount)
        credit_str = "-"
        debit_str = "-"
        if tx.balance_after is not None:
            bal = Decimal(tx.balance_after)
            if bal >= running:
                credit_str = f"INR {float(bal - running):,.2f}"
            else:
                debit_str = f"INR {float(running - bal):,.2f}"
            running = bal
        else:
            if tx.transaction_type in ["DEPOSIT", "REVERSAL"]:
                credit_str = f"INR {float(amt):,.2f}"
            else:
                debit_str = f"INR {float(amt):,.2f}"

        tx_table_data.append([
            Paragraph(tx.created_at.strftime("%Y-%m-%d %H:%M") if tx.created_at else "", cell_style),
            Paragraph(tx.reference_number, cell_style),
            Paragraph(tx.narration[:30], cell_style),
            Paragraph(tx.transaction_mode or "TRANSFER", cell_style),
            Paragraph(credit_str, cell_style),
            Paragraph(debit_str, cell_style),
            Paragraph(f"INR {float(tx.balance_after or 0):,.2f}", cell_style),
            Paragraph(tx.status, cell_style)
        ])

    tx_table = Table(tx_table_data, colWidths=[70, 75, 110, 55, 60, 60, 60, 50])
    tx_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
    ]))
    elements.append(tx_table)

    elements.append(Spacer(1, 20))
    elements.append(Paragraph("<i>This is a computer-generated account statement from BankVCS 2.0 and does not require a physical signature.</i>", sub_style))

    doc.build(elements)
    return buf.getvalue()


def export_transaction_receipt_pdf(transaction_id: int) -> bytes:
    """Generates official transaction receipt PDF byte buffer for a single transaction."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    tx = db.session.get(Transaction, transaction_id)
    if not tx:
        return b""

    account = db.session.get(Account, tx.account_id)
    owner = account.owner if account else None
    owner_name = owner.full_name or owner.username if owner else "Valued Customer"

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'ReceiptTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#0F172A"),
        alignment=1
    )
    sub_style = ParagraphStyle(
        'ReceiptSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748B"),
        alignment=1
    )
    lbl_style = ParagraphStyle(
        'LblStyle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#334155")
    )
    val_style = ParagraphStyle(
        'ValStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#0F172A")
    )

    elements = []

    elements.append(Paragraph("<b>BankVCS 2.0</b>", title_style))
    elements.append(Paragraph("OFFICIAL PAYMENT &amp; TRANSACTION RECEIPT", sub_style))
    elements.append(Spacer(1, 10))
    elements.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor("#10B981"), spaceAfter=15))

    receipt_rows = [
        [Paragraph("Transaction ID / Ref:", lbl_style), Paragraph(tx.reference_number, val_style)],
        [Paragraph("Date &amp; Time:", lbl_style), Paragraph(tx.created_at.strftime("%Y-%m-%d %H:%M:%S UTC") if tx.created_at else "N/A", val_style)],
        [Paragraph("Account Holder:", lbl_style), Paragraph(owner_name, val_style)],
        [Paragraph("Account Number:", lbl_style), Paragraph(mask_account_number(account.account_number if account else ""), val_style)],
        [Paragraph("Transaction Type:", lbl_style), Paragraph(tx.transaction_type, val_style)],
        [Paragraph("Transaction Mode:", lbl_style), Paragraph(tx.transaction_mode or "TRANSFER", val_style)],
        [Paragraph("Amount:", lbl_style), Paragraph(f"<b>INR {float(tx.amount):,.2f}</b>", val_style)],
        [Paragraph("Status:", lbl_style), Paragraph(f"<b>{tx.status}</b>", val_style)],
        [Paragraph("Narration / Remarks:", lbl_style), Paragraph(tx.narration, val_style)],
        [Paragraph("Balance After Transaction:", lbl_style), Paragraph(f"INR {float(tx.balance_after or 0):,.2f}", val_style)],
    ]

    tbl = Table(receipt_rows, colWidths=[180, 324])
    tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#CBD5E1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 12),
        ('RIGHTPADDING', (0, 0), (-1, -1), 12),
    ]))

    elements.append(tbl)
    elements.append(Spacer(1, 25))
    elements.append(Paragraph("<i>Security Notice: This is a verified, computer-generated transaction receipt. No physical signature is required. BankVCS 2.0 Version Control System.</i>", sub_style))

    doc.build(elements)
    return buf.getvalue()
